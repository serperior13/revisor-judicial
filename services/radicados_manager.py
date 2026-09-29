# services/radicados_manager.py
"""
Gestor de listas de radicados (RAMA / SAMAI).

- Lee y escribe los Excel de entrada (Procesos.xlsx, Radicados_SAMAI.xlsx).
- Lectura FLEXIBLE: detecta la hoja con datos y la columna del radicado.
- Al guardar cambios, respalda el Excel anterior en respaldos_radicados/.
"""

import logging
import os
import re
import shutil
from datetime import datetime
from pathlib import Path

import openpyxl

from core.constantes import (
    ARCHIVO_EXCEL,
    ARCHIVO_ENTRADA_SAMAI,
    CARPETA_RESPALDOS_RADICADOS,
    COLUMNA_RADICADO,
    COLUMNA_RADICADO_SAMAI,
)

log = logging.getLogger(__name__)


# ====================================================================
# CONFIGURACIÓN POR JURISDICCIÓN
# ====================================================================

CONFIG_JURISDICCIONES = {
    "RAMA": {
        "archivo": ARCHIVO_EXCEL,
        "nombre_hoja": "Procesos",
        "columna": COLUMNA_RADICADO,
        "prefijo_backup": "Procesos",
    },
    "SAMAI": {
        "archivo": ARCHIVO_ENTRADA_SAMAI,
        "nombre_hoja": "SAMAI",
        "columna": COLUMNA_RADICADO_SAMAI,
        "prefijo_backup": "Radicados_SAMAI",
    },
}

# Nombres de hoja habituales (para intentar primero)
HOJAS_PREFERIDAS = ["Procesos", "SAMAI", "Hoja1", "Sheet1", "Hoja 1", "Datos"]


# ====================================================================
# VALIDACIÓN
# ====================================================================

def limpiar_radicado(texto: str) -> str:
    """Quita todo excepto dígitos."""
    return re.sub(r"\D", "", str(texto or "")).strip()


def validar_radicado(texto: str, jurisdiccion: str) -> tuple[bool, str]:
    """Valida un radicado. Devuelve (es_valido, mensaje_error)."""
    limpio = limpiar_radicado(texto)
    if not limpio:
        return False, "El radicado no puede estar vacío."
    if len(limpio) < 20:
        return False, (
            "El radicado debe tener al menos 20 dígitos.\n"
            "Ingresaste {} dígitos.".format(len(limpio))
        )
    if len(limpio) > 25:
        return False, (
            "El radicado tiene demasiados dígitos ({}). Máximo 25.".format(
                len(limpio))
        )
    return True, ""


# ====================================================================
# LECTURA FLEXIBLE
# ====================================================================

def _es_celda_vacia(valor) -> bool:
    if valor is None:
        return True
    texto = str(valor).strip()
    if not texto:
        return True
    if "borra esta fila" in texto.lower():
        return True
    return False


def _detectar_columna_radicado(hoja, columna_preferida: str) -> int:
    """
    Busca la columna cuyo encabezado contenga 'radica'.
    Si no la encuentra, devuelve 1 (columna A).
    """
    try:
        for fila in hoja.iter_rows(min_row=1, max_row=1):
            for i, celda in enumerate(fila, start=1):
                if celda.value is None:
                    continue
                texto = str(celda.value).lower()
                if "radica" in texto:
                    return i
    except Exception:
        pass

    # Fallback: probar por el nombre exacto de la columna preferida
    try:
        for fila in hoja.iter_rows(min_row=1, max_row=1):
            for i, celda in enumerate(fila, start=1):
                if celda.value is None:
                    continue
                if str(celda.value).strip().lower() == columna_preferida.lower():
                    return i
    except Exception:
        pass

    return 1


def _leer_excel_flexible(ruta_archivo: str, columna_preferida: str,
                         nombre_hoja_preferido: str) -> list[str]:
    """
    Lee un Excel de radicados de forma tolerante:
      1. Intenta la hoja con nombre preferido.
      2. Si no existe, busca entre HOJAS_PREFERIDAS.
      3. Si tampoco, usa la hoja activa.
      4. Detecta la columna del radicado por encabezado.
      5. Lee desde la fila 2.
    """
    if not os.path.exists(ruta_archivo):
        return []

    wb = openpyxl.load_workbook(ruta_archivo, data_only=True, read_only=True)

    # 1) Elegir hoja
    hoja = None
    if nombre_hoja_preferido in wb.sheetnames:
        hoja = wb[nombre_hoja_preferido]
    else:
        for nombre in HOJAS_PREFERIDAS:
            if nombre in wb.sheetnames:
                hoja = wb[nombre]
                break
    if hoja is None:
        hoja = wb.active
    if hoja is None:
        wb.close()
        return []

    # 2) Detectar columna
    idx_col = _detectar_columna_radicado(hoja, columna_preferida)

    # 3) Leer radicados desde fila 2
    radicados = []
    try:
        for fila in hoja.iter_rows(min_row=2, min_col=idx_col,
                                   max_col=idx_col, values_only=True):
            valor = fila[0] if fila else None
            if _es_celda_vacia(valor):
                continue
            limpio = limpiar_radicado(valor)
            if limpio and limpio not in radicados:
                radicados.append(limpio)
    except Exception as e:
        log.warning("Error leyendo filas de %s: %s", ruta_archivo, e)

    wb.close()
    return radicados


# ====================================================================
# GESTOR
# ====================================================================

class RadicadosManager:
    """Gestiona las listas de radicados por jurisdicción."""

    def __init__(self):
        self._cache = {}

    # ------------------------------------------------------------------
    # API pública
    # ------------------------------------------------------------------

    def cargar(self, jurisdiccion: str) -> list[str]:
        jurisdiccion = jurisdiccion.upper()
        cfg = CONFIG_JURISDICCIONES.get(jurisdiccion)
        if not cfg:
            return []

        archivo = cfg["archivo"]
        try:
            radicados = _leer_excel_flexible(
                archivo,
                columna_preferida=cfg["columna"],
                nombre_hoja_preferido=cfg["nombre_hoja"],
            )
            self._cache[jurisdiccion] = radicados
            log.info("📖 %s: %s radicados cargados.", jurisdiccion, len(radicados))
            return radicados
        except Exception as e:
            log.exception("Error leyendo %s: %s", archivo, e)
            return []

    def guardar(self, jurisdiccion: str, radicados: list[str]) -> tuple[bool, str]:
        jurisdiccion = jurisdiccion.upper()
        cfg = CONFIG_JURISDICCIONES.get(jurisdiccion)
        if not cfg:
            return False, "Jurisdicción desconocida."

        archivo = cfg["archivo"]

        # 1) Respaldar el anterior
        respaldo_hecho = False
        if os.path.exists(archivo):
            try:
                ts = datetime.now().strftime("%Y-%m-%d_%H-%M-%S")
                nombre_respaldo = "{}_{}.xlsx".format(cfg["prefijo_backup"], ts)
                ruta_respaldo = os.path.join(
                    CARPETA_RESPALDOS_RADICADOS, nombre_respaldo)
                shutil.copy2(archivo, ruta_respaldo)
                respaldo_hecho = True
                log.info("📦 Respaldo creado: %s", ruta_respaldo)
            except Exception as e:
                log.warning("No se pudo respaldar el Excel anterior: %s", e)

        # 2) Escribir el nuevo Excel
        try:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = cfg["nombre_hoja"]

            from openpyxl.styles import Font, PatternFill, Alignment
            ws["A1"] = cfg["columna"]
            ws["A1"].font = Font(bold=True, color="FFFFFF", size=12)
            ws["A1"].fill = PatternFill(
                start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
            ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
            ws.row_dimensions[1].height = 28
            ws.column_dimensions["A"].width = 42

            for row in range(1, 5001):
                ws.cell(row=row, column=1).number_format = "@"

            for i, rad in enumerate(radicados, start=2):
                ws.cell(row=i, column=1, value=rad).alignment = Alignment(
                    horizontal="left", vertical="center")

            wb.save(archivo)
            self._cache[jurisdiccion] = list(radicados)
            log.info("✅ %s: %s radicados guardados.",
                     jurisdiccion, len(radicados))

            mensaje = "Guardados {} radicados.".format(len(radicados))
            if respaldo_hecho:
                mensaje += " Se respaldó la versión anterior."
            return True, mensaje

        except Exception as e:
            log.exception("Error guardando %s: %s", archivo, e)
            return False, "No se pudo guardar: {}".format(e)

    def agregar(self, jurisdiccion: str, texto: str) -> tuple[bool, str]:
        jurisdiccion = jurisdiccion.upper()
        valido, error = validar_radicado(texto, jurisdiccion)
        if not valido:
            return False, error

        limpio = limpiar_radicado(texto)
        lista = self._cache.get(jurisdiccion)
        if lista is None:
            lista = self.cargar(jurisdiccion)

        if limpio in lista:
            return False, "Ese radicado ya está en la lista."

        lista.append(limpio)
        self._cache[jurisdiccion] = lista
        return True, ""

    def eliminar(self, jurisdiccion: str, radicado: str) -> bool:
        jurisdiccion = jurisdiccion.upper()
        lista = self._cache.get(jurisdiccion, [])
        try:
            lista.remove(radicado)
            self._cache[jurisdiccion] = lista
            return True
        except ValueError:
            return False

    def listar_actuales(self, jurisdiccion: str) -> list[str]:
        jurisdiccion = jurisdiccion.upper()
        if jurisdiccion not in self._cache:
            return self.cargar(jurisdiccion)
        return list(self._cache[jurisdiccion])

    def descartar_cambios(self, jurisdiccion: str) -> None:
        jurisdiccion = jurisdiccion.upper()
        self._cache.pop(jurisdiccion, None)

    # ------------------------------------------------------------------
    # Respaldos
    # ------------------------------------------------------------------

    def listar_respaldos(self, jurisdiccion: str) -> list[Path]:
        cfg = CONFIG_JURISDICCIONES.get(jurisdiccion.upper())
        if not cfg:
            return []

        prefijo = cfg["prefijo_backup"] + "_"
        try:
            archivos = [
                p for p in Path(CARPETA_RESPALDOS_RADICADOS).iterdir()
                if p.is_file()
                and p.name.startswith(prefijo)
                and p.suffix.lower() == ".xlsx"
            ]
            archivos.sort(key=lambda p: p.stat().st_mtime, reverse=True)
            return archivos
        except Exception as e:
            log.warning("Error listando respaldos: %s", e)
            return []

    def restaurar_respaldo(self, jurisdiccion: str, archivo: Path) -> tuple[bool, str]:
        jurisdiccion = jurisdiccion.upper()
        cfg = CONFIG_JURISDICCIONES.get(jurisdiccion)
        if not cfg:
            return False, "Jurisdicción desconocida."

        if not archivo.exists():
            return False, "El archivo de respaldo no existe."

        try:
            archivo_actual = cfg["archivo"]
            if os.path.exists(archivo_actual):
                ts = datetime.now().strftime("%Y-%m-%d_%H-%M-%S")
                nombre_respaldo = "{}_{}.xlsx".format(cfg["prefijo_backup"], ts)
                ruta_respaldo = os.path.join(
                    CARPETA_RESPALDOS_RADICADOS, nombre_respaldo)
                shutil.copy2(archivo_actual, ruta_respaldo)

            shutil.copy2(archivo, archivo_actual)
            self.descartar_cambios(jurisdiccion)
            self.cargar(jurisdiccion)
            return True, "Respaldo restaurado correctamente."

        except Exception as e:
            log.exception("Error restaurando respaldo: %s", e)
            return False, "No se pudo restaurar: {}".format(e)