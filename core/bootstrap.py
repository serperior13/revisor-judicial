# core/bootstrap.py
"""
Crea los archivos de entrada necesarios si no existen.
Se llama al iniciar la aplicación (main.py).

Archivos:
  - Procesos.xlsx           → Radicados de la Rama Judicial
  - Radicados_SAMAI.xlsx    → Radicados del Consejo de Estado (SAMAI)
"""

import logging
import os

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

from core.constantes import (
    ARCHIVO_EXCEL,
    ARCHIVO_ENTRADA_SAMAI,
    COLUMNA_RADICADO,
    COLUMNA_RADICADO_SAMAI,
)

log = logging.getLogger(__name__)


def _crear_excel_plantilla(ruta, nombre_hoja, nombre_columna, ejemplo):
    """Crea un Excel de una sola columna en formato texto."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = nombre_hoja

    # --- Encabezado ---
    ws["A1"] = nombre_columna
    ws["A1"].font = Font(bold=True, color="FFFFFF", size=12)
    ws["A1"].fill = PatternFill(
        start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28
    ws.column_dimensions["A"].width = 42

    # --- Fila de ejemplo ---
    ws["A2"] = ejemplo
    ws["A2"].alignment = Alignment(horizontal="left", vertical="center")
    ws["A2"].font = Font(italic=True, color="808080")

    # --- Aplicar formato TEXTO a toda la columna A ---
    # Esto evita que Excel convierta radicados largos a notación científica
    for row in range(1, 2001):
        ws.cell(row=row, column=1).number_format = "@"

    wb.save(ruta)
    log.info("📄 Plantilla creada: %s", ruta)


def asegurar_archivos_entrada():
    """Crea los dos Excel de entrada si no existen."""
    # --- Procesos.xlsx (Rama Judicial) ---
    if not os.path.exists(ARCHIVO_EXCEL):
        try:
            _crear_excel_plantilla(
                ARCHIVO_EXCEL,
                "Procesos",
                COLUMNA_RADICADO,
                "Borra esta fila y coloca aquí tus radicados (Rama Judicial)",
            )
        except Exception as e:
            log.error("No se pudo crear %s: %s", ARCHIVO_EXCEL, e)

    # --- Radicados_SAMAI.xlsx ---
    if not os.path.exists(ARCHIVO_ENTRADA_SAMAI):
        try:
            _crear_excel_plantilla(
                ARCHIVO_ENTRADA_SAMAI,
                "SAMAI",
                COLUMNA_RADICADO_SAMAI,
                "Borra esta fila y coloca aquí tus radicados (SAMAI)",
            )
        except Exception as e:
            log.error("No se pudo crear %s: %s", ARCHIVO_ENTRADA_SAMAI, e)