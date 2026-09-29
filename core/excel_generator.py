# core/excel_generator.py
"""
Generación del Excel multi-hoja del informe final.
Funciona tanto para Rama Judicial como para SAMAI.
"""

import logging
from datetime import datetime, timedelta

import openpyxl
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils.dataframe import dataframe_to_rows

from core.constantes import (
    ALTO_ENCABEZADO_PUNTOS,
    ALTO_FILA_PUNTOS,
    ANCHOS_COLUMNAS,
    COLUMNA_LINEAS_INTERNA,
)

log = logging.getLogger(__name__)


# ====================================================================
# COLORES (mismos para RAMA y SAMAI)
# ====================================================================

FILL_REVISAR = PatternFill(
    start_color="ED7D31", end_color="ED7D31", fill_type="solid")       # Naranja
FILL_FUERA = PatternFill(
    start_color="FFE699", end_color="FFE699", fill_type="solid")       # Amarillo
FILL_ERROR = PatternFill(
    start_color="FF0000", end_color="FF0000", fill_type="solid")       # Rojo
FILL_SIN_FECHA = PatternFill(
    start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")       # Gris

FONT_BLANCO_BOLD = Font(bold=True, color="FFFFFF")


# ====================================================================
# CLASIFICACIÓN
# ====================================================================

def _clasificar(estado, fecha_str, dias_umbral):
    """
    Determina la clasificación de un registro:
      - 'error'        → falló la consulta
      - 'sin_resultado'→ sin fecha de actuación
      - 'revisar'      → última actuación dentro del umbral
      - 'fuera'        → última actuación fuera del umbral
    """
    estado = str(estado or "")
    if "⚠" in estado or "Error" in estado or "no generó" in estado.lower() \
            or "sin resultado" in estado.lower():
        return "error"
    if not fecha_str:
        return "sin_resultado"
    try:
        f = datetime.strptime(str(fecha_str)[:10], "%Y-%m-%d")
    except (ValueError, TypeError):
        return "sin_resultado"
    if f >= datetime.now() - timedelta(days=dias_umbral):
        return "revisar"
    return "fuera"


def _aplicar_color(cell, clasif):
    """Aplica el color de fondo según la clasificación."""
    if clasif == "revisar":
        cell.fill = FILL_REVISAR
        cell.font = FONT_BLANCO_BOLD
    elif clasif == "fuera":
        cell.fill = FILL_FUERA
    elif clasif == "error":
        cell.fill = FILL_ERROR
        cell.font = FONT_BLANCO_BOLD
    elif clasif == "sin_resultado":
        cell.fill = FILL_SIN_FECHA


# ====================================================================
# GENERADOR PRINCIPAL
# ====================================================================

def generar_excel_multihoja(df_resultados: pd.DataFrame, filepath: str,
                            dias_umbral: int = 10) -> None:
    """
    Genera el Excel con varias hojas y formato condicional.

    Hojas creadas:
      - General
      - Para Revisar
      - Errores y Caídas
      - Fuera de Búsqueda
    """
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    # Normalizar columnas que puedan faltar
    for col in ("Estado", "Fecha_Ultima_Actuacion"):
        if col not in df_resultados.columns:
            df_resultados[col] = ""

    # Clasificación previa (para filtrar hojas)
    clasifs = [
        _clasificar(row.get("Estado", ""),
                    row.get("Fecha_Ultima_Actuacion", ""),
                    dias_umbral)
        for _, row in df_resultados.iterrows()
    ]
    df_resultados = df_resultados.copy()
    df_resultados["_clasif"] = clasifs

    df_general = df_resultados.copy()
    df_revisar = df_resultados[df_resultados["_clasif"] == "revisar"].copy()
    df_fallos = df_resultados[df_resultados["_clasif"] == "error"].copy()
    df_fuera = df_resultados[df_resultados["_clasif"] == "fuera"].copy()

    hojas = [
        ("General", df_general),
        ("Para Revisar", df_revisar),
        ("Errores y Caídas", df_fallos),
        ("Fuera de Búsqueda", df_fuera),
    ]

    alineacion_cuerpo = Alignment(
        wrap_text=True, vertical="center", horizontal="left")
    alineacion_header = Alignment(
        wrap_text=True, vertical="center", horizontal="center")

    for nombre_hoja, df_sub in hojas:
        ws = wb.create_sheet(title=nombre_hoja)

        # Columnas internas no se escriben
        df_visible = df_sub.drop(
            columns=[COLUMNA_LINEAS_INTERNA, "_clasif"], errors="ignore")

        for row in dataframe_to_rows(df_visible, index=False, header=True):
            ws.append(row)

        # Anchos de columna
        for letra, ancho in ANCHOS_COLUMNAS.items():
            ws.column_dimensions[letra].width = ancho

        # Encabezados
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.alignment = alineacion_header
        ws.row_dimensions[1].height = ALTO_ENCABEZADO_PUNTOS

        # Filas de datos
        for row_idx in range(2, ws.max_row + 1):
            for cell in ws[row_idx]:
                cell.alignment = alineacion_cuerpo
            ws.row_dimensions[row_idx].height = ALTO_FILA_PUNTOS

        # Colorear según _clasif
        col_estado_idx = None
        for col in range(1, ws.max_column + 1):
            if ws.cell(row=1, column=col).value == "Estado":
                col_estado_idx = col
                break

        if col_estado_idx and "_clasif" in df_sub.columns:
            for i, (_, row) in enumerate(df_sub.iterrows()):
                row_idx = i + 2
                cell = ws.cell(row=row_idx, column=col_estado_idx)
                _aplicar_color(cell, row["_clasif"])

    wb.save(filepath)
    log.info("🎨 Archivo Excel generado: %s", filepath)


def preparar_dataframe(resultados: list, columnas: list) -> pd.DataFrame:
    """
    Toma una lista de dicts (o RegistroResultado) y devuelve un DataFrame
    con las columnas en el orden correcto, rellenando faltantes.
    """
    filas = []
    for r in resultados:
        if hasattr(r, "to_dict"):
            filas.append(r.to_dict())
        else:
            filas.append(r)

    df = pd.DataFrame(filas)
    for c in columnas:
        if c not in df.columns:
            df[c] = [] if c == COLUMNA_LINEAS_INTERNA else ""
    return df.reindex(columns=columnas)